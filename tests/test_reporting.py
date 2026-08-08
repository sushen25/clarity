from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from app.clinical import CRITERIA
from app.reporting import generate_docx


def test_generated_report_contains_all_clinician_criteria(tmp_path):
    output = tmp_path / "criteria-report.docx"
    outcomes = ("met", "not_met", "insufficient", "unreviewed")
    criteria = [
        {"criterion_id": identifier, "clinician_outcome": outcomes[index % len(outcomes)]}
        for index, identifier in enumerate(CRITERIA)
    ]
    generate_docx(
        template_path=Path(__file__).parents[1] / "templates" / "adhd_report_template.docx",
        output_path=output,
        case={
            "patient_initials": "S. T.", "cohort": "adult",
            "assessment_dates": ["2026-08-08"], "instruments": ["Synthetic instrument"],
        },
        draft={
            "sections": [
                {
                    "key": "background", "heading": "Background",
                    "paragraphs": [{"text": "Verified clinical narrative.", "evidence_ids": ["evidence-1"]}],
                },
            ],
            "version": 1, "model_id": "synthetic-model", "prompt_version": "test-v1", "state": "draft",
        },
        clinician={"full_name": "Dr Synthetic Clinician", "registration_number": "TEST0001"},
        criteria=criteria,
    )

    report = Document(output)
    report_text = "\n".join(paragraph.text for paragraph in report.paragraphs)
    assert "Verified clinical narrative." in report_text
    assert "Evidence:" not in report_text
    assert "evidence-1" not in report_text
    assert report.settings.odd_and_even_pages_header_footer is True
    assert report.sections[0].different_first_page_header_footer is True
    assert report.sections[0].even_page_header.paragraphs[0].text == report.sections[0].header.paragraphs[0].text
    assert report.sections[0].even_page_footer.paragraphs[0].text == report.sections[0].footer.paragraphs[0].text
    assert report.sections[0].first_page_header.paragraphs[0].text == report.sections[0].header.paragraphs[0].text
    assert report.sections[0].first_page_footer.paragraphs[0].text == report.sections[0].footer.paragraphs[0].text
    section_props = report.sections[0]._sectPr
    header_relationships = {
        reference.get(qn("r:id")) for reference in section_props.findall(qn("w:headerReference"))
    }
    footer_relationships = {
        reference.get(qn("r:id")) for reference in section_props.findall(qn("w:footerReference"))
    }
    assert len(header_relationships) == 1
    assert len(footer_relationships) == 1
    criteria_tables = [
        table for table in report.tables
        if table.rows and table.cell(0, 0).text in {
            "A1. Inattention Criteria", "A2. Hyperactivity / Impulsivity Criteria",
        }
    ]
    assert len(criteria_tables) == 2
    assert all(len(table.rows) == 10 for table in criteria_tables)
    assert criteria_tables[0].cell(1, 0).text == "a.   Careless mistakes / attention to detail"
    assert criteria_tables[0].cell(1, 1).text == "✓  Met"
    assert criteria_tables[0].cell(2, 1).text == "Not met"
    assert criteria_tables[0].cell(3, 1).text == "Insufficient evidence"
    assert criteria_tables[0].cell(4, 1).text == "Not reviewed"
    assert criteria_tables[1].cell(9, 0).text == "i.   Interrupts or intrudes"

    for table in criteria_tables:
        assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        assert all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None for row in table.rows)
        grid_widths = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid]
        assert grid_widths == [7560, 1944]
