from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .clinical import CRITERIA_LABELS, REPORT_SECTIONS, RECOMMENDATION_GROUPS, outcome_columns


TEAL = "6EB7C3"
SLATE = RGBColor.from_string("173A44")
MUTED = RGBColor.from_string("64767A")
TABLE_WIDTHS_DXA = (7560, 1944)
ADULT_TABLE_WIDTHS_DXA = (5616, 1944, 1944)


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def _heading(doc: Document, text: str, page_break_before: bool = False) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(text.upper())
    run.bold = True
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.6)
    cell = table.cell(0, 0)
    _shade(cell, TEAL)
    cell.text = ""
    cell.height = Pt(4)


def _set_cell_margins(cell, top: int = 100, start: int = 120,
                      bottom: int = 100, end: int = 120) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    props.append(repeat)


def _prevent_row_split(row) -> None:
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:cantSplit"))


def _set_table_geometry(table, widths: tuple[int, ...]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    props = table._tbl.tblPr
    width = props.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        props.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    layout = props.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        props.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for cell, value in zip(row.cells, widths):
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:w"), str(value))
            tc_width.set(qn("w:type"), "dxa")


def _write_cell(cell, text: str, *, bold: bool = False, size: float = 9.5,
                alignment=WD_ALIGN_PARAGRAPH.LEFT, color: RGBColor = SLATE) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial Narrow"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial Narrow")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial Narrow")
    run.font.size = Pt(size)
    run.font.color.rgb = color


def _criterion_status(outcome: str) -> str:
    return {
        "met": "✓  Met",
        "not_met": "Not met",
        "insufficient": "Insufficient evidence",
        "unreviewed": "Not reviewed",
    }.get(outcome, "Not reviewed")


def _criteria_table(doc: Document, criteria: list[dict], prefix: str, title: str, cohort: str) -> None:
    by_id = {item["criterion_id"]: item for item in criteria}
    columns = outcome_columns(cohort)
    table = doc.add_table(rows=1, cols=1 + len(columns))
    table.style = "Table Grid"
    header = table.rows[0]
    _set_repeat_header(header)
    _prevent_row_split(header)
    for cell in header.cells:
        _shade(cell, TEAL)
    _write_cell(header.cells[0], title, bold=True, size=10.5)
    for index, (_, label) in enumerate(columns, 1):
        _write_cell(header.cells[index], label, bold=True, size=9.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    for index in range(1, 10):
        identifier = f"{prefix}.{index}"
        item = by_id.get(identifier, {})
        row = table.add_row()
        _prevent_row_split(row)
        letter = chr(ord("a") + index - 1)
        _write_cell(row.cells[0], f"{letter}.   {CRITERIA_LABELS[identifier]}")
        for column_index, (field, _) in enumerate(columns, 1):
            outcome = item.get(field, "unreviewed")
            status_color = RGBColor.from_string("24633C") if outcome == "met" else MUTED
            _write_cell(row.cells[column_index], _criterion_status(outcome), bold=outcome == "met", size=9,
                        alignment=WD_ALIGN_PARAGRAPH.CENTER, color=status_color)
    _set_table_geometry(table, ADULT_TABLE_WIDTHS_DXA if cohort == "adult" else TABLE_WIDTHS_DXA)


def _criteria_report(doc: Document, criteria: list[dict], cohort: str) -> None:
    _criteria_table(doc, criteria, "A1", "A1. Inattention Criteria", cohort)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    _criteria_table(doc, criteria, "A2", "A2. Hyperactivity / Impulsivity Criteria", cohort)


def _mirror_page_furniture(doc: Document) -> None:
    """Use a single default page style so Word and LibreOffice repeat furniture.

    Sharing one relationship among first/even/default headers causes some
    LibreOffice versions to lose the header and top margin on continuation pages.
    """
    doc.settings.odd_and_even_pages_header_footer = False
    for section in doc.sections:
        section.different_first_page_header_footer = False
        for reference_name in ("headerReference", "footerReference"):
            for reference in list(section._sectPr.findall(qn(f"w:{reference_name}"))):
                if reference.get(qn("w:type")) != "default":
                    section._sectPr.remove(reference)


def generate_docx(template_path: Path, output_path: Path, case: dict, draft: dict,
                  clinician: dict, criteria: list[dict]) -> None:
    doc = Document(template_path)
    _mirror_page_furniture(doc)
    subheading = doc.styles["Heading 2"]
    subheading.font.name = "Arial Narrow"
    subheading.font.size = Pt(10.5)
    subheading.font.bold = True
    subheading.font.color.rgb = SLATE
    subheading.paragraph_format.keep_with_next = True
    subheading.paragraph_format.space_before = Pt(8)
    subheading.paragraph_format.space_after = Pt(4)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CONFIDENTIAL ASSESSMENT REPORT")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string("173A44")

    metadata = doc.add_table(rows=0, cols=2)
    metadata.autofit = False
    values = [
        ("NAME / INITIALS", case["patient_initials"]),
        ("COHORT", case["cohort"].title()),
        ("DATES ASSESSED", ", ".join(case.get("assessment_dates", []))),
        ("ASSESSMENT INSTRUMENTS", ", ".join(case.get("instruments", []))),
        ("ASSESSMENT CONDUCTED BY", f"{clinician['full_name']}, Psychologist"),
    ]
    for label, value in values:
        cells = metadata.add_row().cells
        cells[0].width, cells[1].width = Inches(1.7), Inches(4.9)
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = value or "Not provided"
    doc.add_paragraph()

    report_sections = list(draft["sections"])
    if not any(section.get("key") == "diagnostic_criteria" for section in report_sections):
        criteria_heading = dict(REPORT_SECTIONS)["diagnostic_criteria"]
        summary_index = next(
            (index for index, section in enumerate(report_sections) if section.get("key") == "summary"),
            len(report_sections),
        )
        report_sections.insert(
            summary_index,
            {"key": "diagnostic_criteria", "heading": criteria_heading, "paragraphs": []},
        )

    previous_was_criteria = False
    for section in report_sections:
        is_criteria = section["key"] == "diagnostic_criteria"
        starts_new_page = is_criteria or previous_was_criteria
        previous_was_criteria = is_criteria
        _heading(doc, section["heading"], page_break_before=starts_new_page)
        if is_criteria:
            _criteria_report(doc, criteria, case["cohort"])
        if not section.get("paragraphs"):
            if is_criteria:
                continue
            paragraph = doc.add_paragraph("Information not supplied or not yet verified.")
            paragraph.runs[0].italic = True
            paragraph.runs[0].font.color.rgb = RGBColor(100, 100, 100)
            continue
        previous_group = None
        for item in section["paragraphs"]:
            group = item.get("recommendation_group") if section["key"] == "recommendations" else None
            if group and group != previous_group:
                heading = doc.add_paragraph(RECOMMENDATION_GROUPS[group], style="Heading 2")
                heading.paragraph_format.keep_with_next = True
                previous_group = group
            if item.get("kind") == "clinician_conclusion":
                doc.add_paragraph("Clinician's diagnostic impression", style="Heading 2")
            paragraph = doc.add_paragraph(item["text"], style="List Bullet" if group else None)
            if item.get("kind") == "missing_information":
                paragraph.runs[0].italic = True
                paragraph.runs[0].font.color.rgb = MUTED
            paragraph.paragraph_format.space_after = Pt(7)

    _heading(doc, "Clinical Review Record", page_break_before=True)
    doc.add_paragraph(
        "This document was generated as a clinician drafting aid. The diagnostic conclusion and criterion decisions "
        "remain the responsibility of the approving clinician."
    )
    review = doc.add_table(rows=0, cols=2)
    for label, value in [
        ("Clinician", clinician["full_name"]),
        ("Registration", clinician.get("registration_number") or "Not provided"),
        ("Draft version", str(draft["version"])),
        ("Model", draft["model_id"]),
        ("Prompt version", draft["prompt_version"]),
        ("Approval state", draft["state"]),
    ]:
        cells = review.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = value

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def render_preview(docx_path: Path, preview_dir: Path, soffice: str) -> Path | None:
    preview_dir.mkdir(parents=True, exist_ok=True)
    binary = shutil.which(soffice)
    if not binary:
        return None
    profile = preview_dir / ".lo-profile"
    profile.mkdir(exist_ok=True)
    command = [binary, "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf",
               "--outdir", str(preview_dir), str(docx_path)]
    completed = subprocess.run(command, capture_output=True, text=True, timeout=120, check=False)
    pdf = preview_dir / f"{docx_path.stem}.pdf"
    if completed.returncode != 0 or not pdf.exists():
        return None
    return pdf
