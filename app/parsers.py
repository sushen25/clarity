from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class DocumentRejected(ValueError):
    pass


ALLOWED_EXTENSIONS = {".txt", ".docx", ".pdf"}


def validate_and_extract(path: Path, original_filename: str) -> tuple[str, str]:
    suffix = Path(original_filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise DocumentRejected("Only TXT, DOCX, and PDF files are accepted.")
    data = path.read_bytes()
    if suffix == ".txt":
        if b"\x00" in data[:4096]:
            raise DocumentRejected("The TXT file appears to be binary.")
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentRejected("TXT files must use UTF-8 encoding.") from exc
        mime = "text/plain"
    elif suffix == ".docx":
        if not data.startswith(b"PK"):
            raise DocumentRejected("The file is not a valid DOCX package.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as package:
                names = {name.lower() for name in package.namelist()}
                if any(name.endswith("vbaproject.bin") for name in names) or "word/vbaproject.bin" in names:
                    raise DocumentRejected("Macro-enabled documents are not accepted.")
                if "word/document.xml" not in names:
                    raise DocumentRejected("The DOCX package has no document body.")
            doc = Document(io.BytesIO(data))
            blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [" | ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()) for cell in row.cells]
                    if any(cells):
                        blocks.append(" | ".join(cells))
            text = "\n\n".join(blocks)
        except DocumentRejected:
            raise
        except Exception as exc:
            raise DocumentRejected("The DOCX file is malformed or unreadable.") from exc
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        if not data.startswith(b"%PDF"):
            raise DocumentRejected("The file is not a valid PDF.")
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise DocumentRejected("Encrypted PDFs are not accepted.")
            pages = []
            for index, page in enumerate(reader.pages, start=1):
                value = (page.extract_text() or "").strip()
                if value:
                    pages.append(f"[Page {index}]\n{value}")
            text = "\n\n".join(pages)
        except DocumentRejected:
            raise
        except Exception as exc:
            raise DocumentRejected("The PDF is malformed or unreadable.") from exc
        if not text.strip():
            raise DocumentRejected("The PDF contains no extractable text; scanned PDFs need OCR before upload.")
        mime = "application/pdf"
    text = re.sub(r"\r\n?", "\n", text).strip()
    if not text:
        raise DocumentRejected("The document contains no usable text.")
    return text, mime

