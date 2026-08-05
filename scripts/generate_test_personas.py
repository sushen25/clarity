"""Generate deterministic, entirely synthetic ADHD workflow test personas."""
from __future__ import annotations

import json
import shutil
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab import rl_config


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "test-data" / "personas"
INK = "173A44"
BLUE = "2E74B5"
MUTED = "66737A"
LIGHT = "E8EEF5"
rl_config.invariant = 1


def src(filename: str, file_type: str, source_type: str, reporter: str, setting: str,
        instrument: str, title: str, sections: list[tuple[str, str]]) -> dict:
    return {
        "filename": filename,
        "file_type": file_type,
        "metadata": {
            "source_type": source_type,
            "reporter": reporter,
            "setting": setting,
            "instrument": instrument,
        },
        "title": title,
        "sections": sections,
    }


PERSONAS = [
    {
        "folder": "01_adult_clear_multisetting",
        "summary": "Adult with longstanding, cross-setting attention and hyperactivity/impulsivity concerns, consistent collateral, and documented differential review.",
        "case": {
            "cohort": "adult",
            "patient_initials": "SK",
            "demographics": {"age": 33, "pronouns": "they/them", "occupation": "product designer", "living_arrangement": "with partner"},
            "referral_question": "Assess longstanding attention, organisation, restlessness, and impulsivity concerns affecting work and home functioning.",
            "assessment_dates": ["2026-06-10", "2026-06-17"],
            "instruments": ["Synthetic Adult Attention Profile (SAAP)", "Synthetic Executive Function Index (SEFI)"],
            "cloud_consent": True,
        },
        "expected": {
            "scenario": "plausible ADHD combined presentation; clinician must decide",
            "expected_warning_fragments_after_verification": [],
            "suggested_final_conclusion": "The supplied synthetic evidence is designed to support a clinician conclusion of ADHD, combined presentation, while noting mild situational anxiety that does not better explain the developmental and cross-setting pattern.",
            "met": ["A1.1", "A1.2", "A1.3", "A1.4", "A1.5", "A1.6", "A1.7", "A1.8", "A1.9", "A2.1", "A2.2", "A2.3", "A2.5", "A2.6", "A2.8"],
            "not_met": ["A2.4", "A2.7", "A2.9"],
            "insufficient": [],
            "contradictions_to_notice": ["SK reports severe conversational interruption; the partner describes it as frequent but manageable with cues."],
        },
        "sources": [
            src("01_assessment_transcript.txt", "txt", "transcript", "patient", "clinical", "", "Synthetic assessment interview transcript", [
                ("Transcript", "CLINICIAN: What led to the referral?\nSK: At work I miss details in design specifications, jump between tasks, and need reminders for deadlines. At home I start chores and leave them half finished. My partner says I interrupt and fidget through meals. These problems have been present for years, not only during stressful weeks.\n\nCLINICIAN: What was childhood like?\nSK: Primary school reports repeatedly said I was capable but rushed, forgot equipment, called out, and needed to sit near the teacher. I lost homework and my parents used a checklist from about age eight. In secondary school I relied on last-minute effort and frequent extensions.\n\nCLINICIAN: Where is the impact now?\nSK: At work I have missed three internal deadlines in six months and a colleague checks my handovers. At home I have paid two bills late and arguments happen when I forget agreed tasks. Socially I interrupt friends and sometimes commit to plans without checking my calendar.\n\nCLINICIAN: Restlessness or impulsivity?\nSK: I tap my feet, leave long meetings for breaks, speak before people finish, and buy hobby equipment impulsively. I can remain seated when required, but it takes effort.\n\nCLINICIAN: Mood, anxiety, sleep, substances, or medical issues?\nSK: I get anxious before deadlines but the attention problems occur when calm. Mood is generally steady. I sleep seven hours most nights. I drink alcohol socially and do not use illicit substances. There is no history of mania, psychosis, seizures, head injury, or current safety risk.\n\nCLINICIAN OBSERVATION: SK arrived on time with notes, was cooperative and reflective, shifted position often, and occasionally answered before a question was complete. Speech was coherent. No acute distress or safety concern was observed.\n\nCLINICIAN: What goes well?\nSK: I am creative, good at visual problem solving, warm with clients, and persistent when a project is meaningful. Timers, written steps, and quiet rooms help."),
            ]),
            src("02_self_history.docx", "docx", "questionnaire_narrative", "patient", "home", "", "Synthetic self-history narrative", [
                ("Developmental and education history", "SK recalls being distractible and disorganised before age 12. Family kept a basket at the front door because school materials were frequently lost. Teachers described rushed work, avoidable errors, talking, and inconsistent completion despite strong reasoning skills. No developmental milestone delays were reported."),
                ("Current daily functioning", "At home, cooking is sometimes interrupted by another task, household paperwork accumulates, and multi-step chores require written lists. At work, notifications derail focus, estimates are overly optimistic, and detailed checking is difficult. In friendships, SK sometimes interrupts or forgets to reply."),
                ("Health and differential information", "A recent general-practice review found no reported neurological illness, untreated thyroid concern, or medication effect likely to explain the pattern. Mild deadline-related anxiety is present, but symptoms also occur during holidays and enjoyable low-stress tasks. Sleep is generally regular."),
                ("Strengths and supports", "SK describes creativity, humour, strong visual reasoning, empathy, and willingness to use calendars, body doubling, noise reduction, and task breakdown."),
            ]),
            src("03_partner_collateral.pdf", "pdf", "collateral_report", "partner", "home", "", "Synthetic partner collateral report", [
                ("Relationship context", "The reporter has lived with SK for six years and observes behaviour at home and in social settings."),
                ("Observed attention and organisation", "SK commonly begins several chores, misplaces keys and headphones, forgets verbal requests, and needs shared-calendar reminders. These behaviours occur even during calm periods and have caused late fees and repeated household conflict."),
                ("Observed activity and impulse control", "SK fidgets during television and meals, changes topic quickly, and frequently finishes another person's sentence. Interruption is frequent but often improves after a discreet cue. Impulsive purchases are occasional rather than constant."),
                ("Strengths and alternative explanations", "SK is caring, imaginative, dependable in emergencies, and highly focused on design problems of interest. The reporter has not observed sustained elevated mood, markedly reduced need for sleep, intoxication-related patterns, or a recent abrupt onset."),
            ]),
            src("04_gp_medical_summary.txt", "txt", "medical_summary", "general practitioner", "clinical", "", "Synthetic medical and mental-health summary", [
                ("Summary", "This is a fictional GP summary for software testing. SK reports lifelong attention and organisation difficulties. Physical examination was unremarkable. Sleep is usually seven hours; there is no reported sleep apnoea, seizure disorder, significant head injury, thyroid disorder, or current substance misuse. Mild performance anxiety is noted without panic, major depressive episode, mania, psychosis, or current self-harm risk. Medical assessment does not identify a better primary explanation. Family history includes a first-degree relative with attention difficulties. Strengths include exercise, stable housing, supportive relationships, and engagement with care."),
            ]),
            src("05_scored_instrument_summary.pdf", "pdf", "instrument_export", "patient and partner", "home", "Synthetic Adult Attention Profile (SAAP); Synthetic Executive Function Index (SEFI)", "Fictional already-scored instrument summary", [
                ("Important notice", "These are fictional instruments and synthetic values created only for application testing. No proprietary questions are reproduced and no score should be used clinically."),
                ("SAAP - patient", "Attention/organisation index: T=72. Activity/impulse index: T=68. Response consistency: acceptable. Supplied interpretation: elevations are consistent with the reported concern but are not diagnostic."),
                ("SAAP - partner", "Attention/organisation index: T=69. Activity/impulse index: T=63. Response consistency: acceptable. Supplied interpretation: cross-informant elevation is present."),
                ("SEFI - patient", "Working organisation index: T=70. Task monitoring index: T=67. Supplied interpretation: executive-function concerns warrant integration with interview, history, collateral, and impairment evidence."),
            ]),
        ],
    },
    {
        "folder": "02_adolescent_conflicting_reporters",
        "summary": "Adolescent with strong school impairment, parent/teacher disagreement, developmental evidence, and anxiety considered as a co-occurring factor.",
        "case": {
            "cohort": "adolescent",
            "patient_initials": "LM",
            "demographics": {"age": 15, "pronouns": "she/her", "school_year": 10, "living_arrangement": "shared care between parents"},
            "referral_question": "Assess attention and executive-function concerns associated with declining school completion and family conflict.",
            "assessment_dates": ["2026-05-04", "2026-05-11"],
            "instruments": ["Synthetic Youth Attention Profile (SYAP)", "Synthetic School Function Index (SSFI)"],
            "cloud_consent": True,
        },
        "expected": {
            "scenario": "plausible ADHD predominantly inattentive presentation with conflicting severity reports",
            "expected_warning_fragments_after_verification": [],
            "suggested_final_conclusion": "The synthetic case is designed for a clinician to consider ADHD, predominantly inattentive presentation, with performance anxiety as a possible co-occurring condition rather than a full explanation.",
            "met": ["A1.1", "A1.2", "A1.3", "A1.4", "A1.5", "A1.7", "A1.8", "A1.9", "A2.2"],
            "not_met": ["A1.6", "A2.1", "A2.3", "A2.4", "A2.5", "A2.6", "A2.7", "A2.8", "A2.9"],
            "insufficient": [],
            "contradictions_to_notice": ["Mother reports few difficulties during highly structured weekends; teacher and father describe substantial independent-task impairment."],
        },
        "sources": [
            src("01_adolescent_interview.txt", "txt", "transcript", "adolescent", "clinical", "", "Synthetic adolescent interview transcript", [
                ("Transcript", "CLINICIAN: What has been difficult?\nLM: I understand lessons but miss instructions, forget which portal has homework, and submit work unfinished. In exams I skip parts of questions. I lose my bus card and sports gear.\n\nCLINICIAN: Was this present when you were younger?\nLM: In primary school I was always moved to the front and teachers checked my diary. Dad packed my bag until Year 7. I daydreamed and finished slowly, but marks were okay because there was more help.\n\nCLINICIAN: Does it happen outside school?\nLM: At Dad's house I forget chores and leave things everywhere. Mum writes a schedule and says I am fine when I follow it. With friends I miss details in plans and arrive without what I need.\n\nCLINICIAN: Restlessness or impulsivity?\nLM: I doodle and bounce my leg, but I do not usually leave my seat or interrupt.\n\nCLINICIAN: Mood, worry, sleep, and safety?\nLM: I worry about grades and sometimes avoid opening the school portal because I expect bad news. Worry became worse after grades dropped. I sleep around eight hours. I do not use alcohol or drugs. I have not had thoughts of harming myself.\n\nCLINICIAN OBSERVATION: LM was thoughtful and initially anxious, checked the question twice, lost her place during a longer explanation, and used written notes effectively. No acute safety concern was observed.\n\nCLINICIAN: Strengths?\nLM: I am good at art, supportive with friends, and work hard when instructions are broken into steps."),
            ]),
            src("02_parent_history.docx", "docx", "collateral_report", "mother and father", "home", "", "Synthetic parent developmental history", [
                ("Early development", "Pregnancy, birth, hearing, vision, and developmental milestones were reported as unremarkable. From early primary school LM needed repeated prompts, visual routines, and adult checking to complete morning and homework tasks."),
                ("Mother's observations", "At mother's home, detailed routines and one-to-one reminders mean tasks are usually completed. Mother sees some forgetfulness but questions whether the problem is severe outside school and believes anxiety may account for much of the current presentation."),
                ("Father's observations", "At father's home, where routines are less externally managed, LM forgets chores, leaves school materials behind, underestimates time, and abandons multi-step tasks. Father recalls the same pattern before age 12."),
                ("Family, health, and strengths", "There is no reported neurological disorder, substance use, mania, or major depressive episode. Grade-related worry is present. LM is creative, kind to younger cousins, persistent with drawing, and responds well to calm structure."),
            ]),
            src("03_teacher_collateral.pdf", "pdf", "collateral_report", "English and science teachers", "school", "", "Synthetic teacher collateral report", [
                ("School observations", "Across English and science, LM misses multi-part directions, leaves question sections blank, loses worksheets, and requires prompts to begin independent work. Difficulties have been evident across two school years and are not confined to one subject."),
                ("Functional impact", "Six assignments were late this semester despite demonstrated understanding in discussion. Teachers report that reduced workload, written instructions, and staged deadlines improve completion. Peer relationships are appropriate."),
                ("Activity and behaviour", "LM is quiet, remains seated, and rarely calls out. Leg movement and doodling are observed but do not disrupt others."),
                ("Strengths and contradictions", "LM contributes original ideas, responds well to feedback, and produces strong artwork. Her presentation appears more impaired at school than described by her mother, possibly because school requires more independent organisation."),
            ]),
            src("04_health_and_wellbeing.txt", "txt", "medical_summary", "school wellbeing clinician", "clinical", "", "Synthetic health and differential summary", [
                ("Summary", "This fictional wellbeing summary records normal hearing and vision screening, no known neurological illness, no significant head injury, and no current medication effect. LM reports performance anxiety linked to unfinished school tasks, without panic disorder, mania, psychosis, substance use, or current self-harm risk. Sleep and appetite are broadly stable.\n\nDifferential formulation remains required and should integrate alternative explanations described in the health, learning, emotional, and environmental history.\n\nAvailable history does not show a single-subject skill deficit, although formal learning assessment has not been supplied. Protective strengths include supportive parents, friendships, art participation, and willingness to use planning supports."),
            ]),
            src("05_scored_instrument_summary.pdf", "pdf", "instrument_export", "adolescent, mother, father, and teacher", "school", "Synthetic Youth Attention Profile (SYAP); Synthetic School Function Index (SSFI)", "Fictional already-scored instrument summary", [
                ("Important notice", "These fictional scales and synthetic values exist only for software testing. They contain no proprietary items and cannot support a real diagnosis."),
                ("SYAP profiles", "Adolescent attention index T=71; mother T=58; father T=68; teacher T=73. Activity/impulse indices range T=49-56. Validity indicators were supplied as acceptable."),
                ("SSFI teacher export", "Independent task completion T=74 and materials management T=70. Supplied interpretation: marked school-function concern, to be integrated with interview, developmental history, other settings, and differential assessment."),
            ]),
        ],
    },
    {
        "folder": "03_adult_negative_anxiety_sleep",
        "summary": "Adult with recent attention complaints better aligned with shift-work sleep disruption and anxiety, with negative childhood and collateral evidence.",
        "case": {
            "cohort": "adult",
            "patient_initials": "RB",
            "demographics": {"age": 41, "pronouns": "he/him", "occupation": "registered nurse", "living_arrangement": "with spouse and child"},
            "referral_question": "Clarify whether recent concentration problems reflect ADHD or another explanation.",
            "assessment_dates": ["2026-04-08", "2026-04-15"],
            "instruments": ["Synthetic Adult Attention Profile (SAAP)", "Synthetic Anxiety and Sleep Context Scale (SASCS)"],
            "cloud_consent": True,
        },
        "expected": {
            "scenario": "negative ADHD conclusion; sleep disruption and anxiety are stronger explanations",
            "expected_warning_fragments_after_verification": [],
            "suggested_final_conclusion": "The synthetic evidence is designed to support a clinician conclusion that ADHD criteria are not met because childhood onset and pervasive cross-setting symptoms are not established, while sleep disruption and anxiety track the recent difficulties.",
            "met": [],
            "not_met": ["A1.1", "A1.2", "A1.3", "A1.4", "A1.5", "A1.6", "A1.7", "A1.8", "A1.9", "A2.1", "A2.2", "A2.3", "A2.4", "A2.5", "A2.6", "A2.7", "A2.8", "A2.9"],
            "insufficient": [],
            "contradictions_to_notice": ["RB describes severe concentration problems after night shifts, while spouse and manager report typical functioning after adequate sleep."],
        },
        "sources": [
            src("01_assessment_transcript.txt", "txt", "transcript", "patient", "clinical", "", "Synthetic differential assessment transcript", [
                ("Transcript", "CLINICIAN: When did concentration become a concern?\nRB: About nine months ago after I moved onto rotating night shifts and my father became ill. I reread medication protocols and feel mentally foggy after poor sleep. On annual leave, after several regular nights, my focus mostly returns.\n\nCLINICIAN: What about childhood and study?\nRB: I do not remember attention or behaviour problems. Reports described me as organised and quiet. I completed homework independently, kept track of equipment, and finished nursing training without extensions.\n\nCLINICIAN: Impact across settings?\nRB: The main impact is during night work. At home after a shift I am irritable and forget small requests, but on days off I manage bills, appointments, cooking, and my child's routines.\n\nCLINICIAN: Hyperactivity or impulsivity?\nRB: No longstanding restlessness, interrupting, excessive talking, or impulsive decisions.\n\nCLINICIAN: Mood and worry?\nRB: I worry about making a clinical error and about my father's health. I lie awake after shifts and sleep four to five broken hours. I feel tense but have no history of mania, psychosis, substance misuse, or self-harm thoughts.\n\nCLINICIAN OBSERVATION: RB appeared tired and worried but remained seated, followed the interview sequence, and gave organised examples. No acute safety concern was observed.\n\nCLINICIAN: Strengths?\nRB: I am careful, compassionate, methodical, and willing to address sleep and anxiety."),
            ]),
            src("02_self_context.docx", "docx", "questionnaire_narrative", "patient", "home", "", "Synthetic self-context narrative", [
                ("Timeline", "Concentration problems began in mid-adulthood alongside rotating shifts, restricted sleep, caring stress, and increased worry. RB reports no comparable childhood or adolescent pattern."),
                ("Function by context", "Errors and rereading occur after night duty. With adequate sleep, RB completes household administration, follows recipes, keeps appointments, and sustains hobbies without notable organisation problems."),
                ("Health and mental health", "Sleep is fragmented around night shifts. Anxiety centres on work safety and a family member's health. RB reports no neurological illness, head injury, mania, psychosis, substance misuse, or current safety risk."),
                ("Strengths", "RB identifies careful practice, empathy, stable relationships, exercise, insight, and readiness to seek sleep and anxiety support."),
            ]),
            src("03_spouse_and_manager_collateral.pdf", "pdf", "collateral_report", "spouse and nurse unit manager", "home and work", "", "Synthetic spouse and workplace collateral", [
                ("Spouse report", "Before rotating shifts, RB was consistently organised with bills, appointments, and parenting tasks. Forgetfulness appears after poor sleep and improves over several rested days. The spouse has not observed lifelong impulsivity or restlessness."),
                ("Manager report", "RB's historical performance was careful and reliable. Recent slowing and repeated checking cluster after night shifts. No pattern of lost equipment, missed handovers, disruptive activity, or impulsive behaviour is reported."),
                ("Strengths and risk", "Both reporters describe RB as conscientious, compassionate, and receptive to feedback. Neither reports intoxication, unsafe impulsivity, or acute mental-health risk."),
            ]),
            src("04_gp_sleep_summary.txt", "txt", "medical_summary", "general practitioner", "clinical", "", "Synthetic medical and sleep summary", [
                ("Summary", "This fictional medical summary records rotating shift-work sleep disruption and anxiety symptoms. Screening blood work is described as unremarkable for software-test purposes. No seizure disorder, significant head injury, untreated thyroid condition, substance misuse, mania, psychosis, or current self-harm risk is reported. The temporal relationship between sleep loss, worry, and concentration should be assessed before attributing symptoms to ADHD. Protective factors include family support, employment, insight, and engagement in treatment."),
            ]),
            src("05_scored_instrument_summary.pdf", "pdf", "instrument_export", "patient and spouse", "home", "Synthetic Adult Attention Profile (SAAP); Synthetic Anxiety and Sleep Context Scale (SASCS)", "Fictional already-scored instrument summary", [
                ("Important notice", "Fictional scales and synthetic values for software testing only; no proprietary items or clinical scoring are included."),
                ("SAAP", "Patient attention index T=64; spouse attention index T=52; activity/impulse indices T=47 and T=45. Supplied interpretation: limited cross-informant convergence and no developmental inference."),
                ("SASCS", "Sleep disruption context index T=76 and worry context index T=70. Supplied interpretation: current concentration complaints strongly covary with sleep and worry; integrate with clinical history."),
            ]),
        ],
    },
    {
        "folder": "04_adolescent_intentionally_incomplete",
        "summary": "Adolescent self-report-only case designed to trigger missing collateral, setting, developmental, medical, differential, and strengths warnings.",
        "case": {
            "cohort": "adolescent",
            "patient_initials": "TN",
            "demographics": {"age": 14, "pronouns": "they/them", "school_year": 9},
            "referral_question": "Explore self-reported concentration problems. Collateral and health history have not yet been obtained.",
            "assessment_dates": ["2026-07-02"],
            "instruments": ["Synthetic Youth Attention Profile (SYAP)"],
            "cloud_consent": True,
        },
        "expected": {
            "scenario": "insufficient evidence; gap-detection fixture",
            "expected_warning_fragments_after_verification": ["Developmental history", "Medical history", "Differential", "strengths", "two important settings", "parent/carer or teacher collateral"],
            "suggested_final_conclusion": "Insufficient information is available to determine whether ADHD criteria are met. Developmental, medical, collateral, differential, and cross-setting evidence should be obtained.",
            "met": [],
            "not_met": [],
            "insufficient": ["A1.1", "A1.2", "A1.3", "A1.4", "A1.5", "A1.6", "A1.7", "A1.8", "A1.9", "A2.1", "A2.2", "A2.3", "A2.4", "A2.5", "A2.6", "A2.7", "A2.8", "A2.9"],
            "contradictions_to_notice": [],
        },
        "sources": [
            src("01_brief_interview.txt", "txt", "transcript", "adolescent", "clinical", "", "Synthetic incomplete interview", [
                ("Transcript", "CLINICIAN: What concerns you?\nTN: I cannot focus in maths and I forget assignments. It feels worse this year.\n\nCLINICIAN: Was this present in primary school or at home?\nTN: I am not sure. Nobody else has been asked yet.\n\nCLINICIAN: Mood, sleep, health, or other possible explanations?\nTN: We did not cover that today.\n\nCLINICIAN OBSERVATION: A brief telehealth interview was completed. The available session was too short for developmental, medical, mental-health, differential, impairment, strengths, or safety assessment."),
            ]),
            src("02_self_report.docx", "docx", "questionnaire_narrative", "adolescent", "school", "", "Synthetic limited self-report", [
                ("Current concern", "TN reports difficulty focusing in mathematics, forgetting some assignments, and feeling overwhelmed by the school portal during the current year."),
                ("Information not yet supplied", "No parent, carer, teacher, medical, developmental, family, mental-health, learning, functional-impairment, or strengths history has been provided. No second setting has been assessed."),
            ]),
            src("03_scored_instrument_summary.pdf", "pdf", "instrument_export", "adolescent", "school", "Synthetic Youth Attention Profile (SYAP)", "Fictional single-respondent instrument summary", [
                ("Important notice", "This fictional instrument and synthetic value are for software testing only. No proprietary questions are reproduced."),
                ("Result", "Adolescent attention index T=71. No validity concern was supplied. No parent, teacher, or other respondent result is available."),
                ("Interpretation", "A single elevated self-report is supporting information only and cannot establish onset, pervasiveness, impairment, differential explanation, or diagnosis."),
            ]),
        ],
    },
    {
        "folder": "05_adult_comorbid_and_conflicting",
        "summary": "Adult with developmental and cross-setting ADHD evidence, clinically significant anxiety, and a university reporter who sees limited outward impairment.",
        "case": {
            "cohort": "adult",
            "patient_initials": "AP",
            "demographics": {"age": 29, "pronouns": "she/her", "occupation": "postgraduate student and casual tutor", "living_arrangement": "share house"},
            "referral_question": "Assess lifelong executive-function and restlessness concerns while considering anxiety and depressive history.",
            "assessment_dates": ["2026-03-03", "2026-03-10", "2026-03-24"],
            "instruments": ["Synthetic Adult Attention Profile (SAAP)", "Synthetic Mood Context Inventory (SMCI)"],
            "cloud_consent": True,
        },
        "expected": {
            "scenario": "plausible ADHD combined presentation with co-occurring anxiety and meaningful masking/compensation",
            "expected_warning_fragments_after_verification": [],
            "suggested_final_conclusion": "The synthetic evidence is designed for a clinician to consider ADHD, combined presentation, alongside a co-occurring anxiety condition. Anxiety increases impairment but does not fully account for the childhood-onset pattern.",
            "met": ["A1.1", "A1.2", "A1.3", "A1.4", "A1.5", "A1.6", "A1.8", "A1.9", "A2.1", "A2.2", "A2.3", "A2.5", "A2.6"],
            "not_met": ["A1.7", "A2.4", "A2.7", "A2.8", "A2.9"],
            "insufficient": [],
            "contradictions_to_notice": ["University supervisor sees punctual, polished work; AP and housemate describe extensive compensatory effort, deadline crises, and home impairment."],
        },
        "sources": [
            src("01_assessment_transcript.txt", "txt", "transcript", "patient", "clinical", "", "Synthetic adult interview with comorbidity", [
                ("Transcript", "CLINICIAN: Describe the current concern.\nAP: I can produce good work, but only through panic, all-night catch-up, many alarms, and rewriting lists. I drift in readings, miss details, underestimate time, and interrupt in seminars. At home I leave appliances on, forget rent transfers, and accumulate unfinished projects.\n\nCLINICIAN: Childhood?\nAP: Before age 12 my reports said I talked, rushed, forgot instructions, and did not show my ability consistently. Mum sat beside me for homework. I was always climbing or moving and was called intense.\n\nCLINICIAN: Other settings?\nAP: My housemate prompts chores and says I dominate conversations. As a tutor I overprepare so students see me as organised. My supervisor sees polished final work but not the missed sleep and extensions.\n\nCLINICIAN: Mental health?\nAP: Anxiety has been present since late adolescence and worsens before evaluation. I had a depressive episode four years ago that is now in remission. Attention, restlessness, and impulsive speech were present well before either problem. No mania, psychosis, substance dependence, or current self-harm thoughts.\n\nCLINICIAN OBSERVATION: AP brought detailed notes, spoke rapidly, interrupted twice, apologised, and returned to the topic with prompts. Affect was anxious but reactive; thought form was coherent. No acute safety concern was observed.\n\nCLINICIAN: Strengths?\nAP: I am curious, energetic, good at teaching, and persistent. Exercise, written agendas, quiet study spaces, and accountability help."),
            ]),
            src("02_developmental_history.docx", "docx", "collateral_report", "mother", "home", "", "Synthetic maternal developmental history", [
                ("Development", "Pregnancy, birth, milestones, hearing, and vision were reported as unremarkable. From early primary school AP needed close homework supervision, lost notices, rushed work, talked frequently, and moved constantly."),
                ("Adolescent functioning", "AP achieved good marks with extensive parental structure, late nights, and repeated checking. Her room and school bag were disorganised, and family arguments occurred around forgotten chores and time management."),
                ("Mental health and health context", "Anxiety became prominent in later adolescence, after attention and activity concerns were already established. Mother reports no childhood seizure, significant head injury, substance problem, manic episode, or psychosis."),
                ("Strengths", "AP is described as warm, imaginative, verbally skilled, resilient, and strongly motivated by helping others."),
            ]),
            src("03_housemate_and_supervisor.pdf", "pdf", "collateral_report", "housemate and university supervisor", "home and university", "", "Synthetic conflicting adult collateral", [
                ("Housemate report", "Across three years, AP has needed reminders for rent and chores, begins many household projects, talks over others when excited, and paces during conversations. Anxiety increases the pattern but it is also visible during relaxed periods."),
                ("Supervisor report", "AP attends supervision punctually with polished work and contributes thoughtful ideas. The supervisor has not observed marked disorganisation or restlessness and considers outward academic functioning strong."),
                ("Context for discrepancy", "AP reports spending disproportionate hours preparing, using multiple reminders, losing sleep near deadlines, and requesting extensions. The supervisor sees final outputs rather than preparation or home functioning."),
                ("Strengths and risk", "Both reporters describe AP as creative, caring, and committed. Neither reports intoxication, mania, psychosis, or acute safety concerns."),
            ]),
            src("04_psychology_and_medical_summary.txt", "txt", "medical_summary", "treating psychologist and general practitioner", "clinical", "", "Synthetic differential and medical summary", [
                ("Summary", "This fictional combined summary notes a history of generalised anxiety symptoms and a past depressive episode now in remission. Attention and restlessness reportedly predate these conditions and persist outside acute mood episodes. Sleep becomes restricted during deadline crises but is otherwise seven to eight hours. No seizure disorder, significant head injury, thyroid concern, substance dependence, mania, psychosis, or current self-harm risk is reported.\n\nFunctional impairment includes late rent transfers, household conflict, repeated extension requests, and disrupted routines.\n\nDifferential review should consider anxiety, mood history, sleep, and compensatory masking. Protective strengths include insight, treatment engagement, exercise, friendships, and meaningful study."),
            ]),
            src("05_scored_instrument_summary.pdf", "pdf", "instrument_export", "patient and mother", "home", "Synthetic Adult Attention Profile (SAAP); Synthetic Mood Context Inventory (SMCI)", "Fictional already-scored instrument summary", [
                ("Important notice", "Fictional instruments and synthetic values for software testing only; no proprietary items are reproduced."),
                ("SAAP", "Patient attention index T=73 and activity/impulse index T=69. Maternal retrospective attention index T=67 and activity/impulse index T=65. Supplied consistency indicators: acceptable."),
                ("SMCI", "Current worry context index T=71; current depressive context index T=54. Supplied interpretation: anxiety is clinically relevant and should be integrated without assuming it explains the full developmental pattern."),
            ]),
        ],
    },
]


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    props = table._tbl.tblPr
    width = props.first_child_found_in("w:tblW")
    width.set(qn("w:w"), str(sum(widths_dxa)))
    width.set(qn("w:type"), "dxa")
    indent = props.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        props.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        for cell, value in zip(row.cells, widths_dxa):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(value))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_docx(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_running_furniture(doc: Document, label: str) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(f"SYNTHETIC TEST FIXTURE  |  {label}")
    set_font(run, 8.5, MUTED, True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Not for clinical use - entirely fictional")
    set_font(run, 8, MUTED, False, True)


def write_docx(path: Path, source: dict, persona_label: str) -> None:
    doc = Document()
    style_docx(doc)
    add_running_furniture(doc, persona_label)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(source["title"])
    set_font(run, 24, INK, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Entirely synthetic input document for Clarity POC testing")
    set_font(run, 10.5, MUTED, False, True)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in (("Reporter", source["metadata"]["reporter"]), ("Setting", source["metadata"]["setting"]), ("Source type", source["metadata"]["source_type"])):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, 10)
    set_table_geometry(table, [2700, 6660])
    for heading, body in source["sections"]:
        doc.add_heading(heading, level=1)
        for paragraph_text in body.split("\n\n"):
            doc.add_paragraph(paragraph_text)
    props = doc.core_properties
    props.title = source["title"]
    props.subject = "Synthetic ADHD report automation test fixture"
    props.author = "Clarity Fixture Generator"
    props.keywords = "synthetic,test fixture,no patient data"
    props.comments = "Entirely fictional. Not for clinical use."
    props.created = datetime(2026, 1, 1, tzinfo=UTC)
    props.modified = datetime(2026, 1, 1, tzinfo=UTC)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    normalized = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path) as source_zip, zipfile.ZipFile(normalized, "w") as target_zip:
        for item in sorted(source_zip.infolist(), key=lambda value: value.filename):
            info = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            info.create_system = item.create_system
            target_zip.writestr(info, source_zip.read(item.filename))
    normalized.replace(path)


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("FixtureTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=colors.HexColor(f"#{INK}"), alignment=TA_CENTER, spaceAfter=7),
        "subtitle": ParagraphStyle("FixtureSubtitle", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=9.5, leading=12, textColor=colors.HexColor(f"#{MUTED}"), alignment=TA_CENTER, spaceAfter=14),
        "h1": ParagraphStyle("FixtureH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=12, spaceAfter=6, keepWithNext=True),
        "body": ParagraphStyle("FixtureBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=14, textColor=colors.HexColor("#202A2E"), spaceAfter=7),
        "small": ParagraphStyle("FixtureSmall", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor(f"#{MUTED}")),
    }


def write_pdf(path: Path, source: dict, persona_label: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = pdf_styles()

    def furniture(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawRightString(7.5 * inch, 10.55 * inch, f"SYNTHETIC TEST FIXTURE | {persona_label}")
        canvas.setFont("Helvetica-Oblique", 8)
        canvas.drawCentredString(4.25 * inch, 0.45 * inch, f"Not for clinical use - entirely fictional | Page {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=0.85 * inch, bottomMargin=0.75 * inch, title=source["title"], author="Clarity Fixture Generator", subject="Synthetic ADHD report automation test fixture")
    story = [Paragraph(source["title"], styles["title"]), Paragraph("Entirely synthetic input document for Clarity POC testing", styles["subtitle"])]
    data = [["Reporter", source["metadata"]["reporter"]], ["Setting", source["metadata"]["setting"]], ["Source type", source["metadata"]["source_type"]]]
    table = Table(data, colWidths=[1.35 * inch, 5.15 * inch], hAlign="LEFT")
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{LIGHT}")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B7C3C9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([table, Spacer(1, 8)])
    for heading, body in source["sections"]:
        story.append(Paragraph(heading, styles["h1"]))
        for paragraph_text in body.split("\n\n"):
            story.append(Paragraph(paragraph_text.replace("&", "&amp;"), styles["body"]))
    document.build(story, onFirstPage=furniture, onLaterPages=furniture)


def write_txt(path: Path, source: dict, persona_label: str) -> None:
    lines = [
        "SYNTHETIC TEST FIXTURE - NOT FOR CLINICAL USE",
        f"Persona: {persona_label}",
        f"Title: {source['title']}",
        f"Reporter: {source['metadata']['reporter']}",
        f"Setting: {source['metadata']['setting']}",
        f"Source type: {source['metadata']['source_type']}",
        "",
    ]
    for heading, body in source["sections"]:
        lines.extend([heading.upper(), body, ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def expanded_expected(expected: dict) -> dict:
    outcomes = {}
    for identifier in [f"A1.{i}" for i in range(1, 10)] + [f"A2.{i}" for i in range(1, 10)]:
        if identifier in expected["met"]:
            value = "met"
        elif identifier in expected["not_met"]:
            value = "not_met"
        else:
            value = "insufficient"
        outcomes[identifier] = {"clinician_outcome": value, "notes": "Synthetic expected-review oracle; clinician verification still required."}
    result = dict(expected)
    result["criterion_outcomes"] = outcomes
    for key in ("met", "not_met", "insufficient"):
        result.pop(key)
    result["do_not_upload"] = True
    return result


def main() -> None:
    if OUTPUT.exists():
        shutil.rmtree(OUTPUT)
    OUTPUT.mkdir(parents=True)
    catalog = []
    manifest_paths = []
    for persona in PERSONAS:
        folder = OUTPUT / persona["folder"]
        inputs = folder / "inputs"
        expected_dir = folder / "expected"
        inputs.mkdir(parents=True)
        expected_dir.mkdir(parents=True)
        (folder / "case.json").write_text(json.dumps(persona["case"], indent=2) + "\n", encoding="utf-8")
        source_manifest = []
        for order, source in enumerate(persona["sources"], start=1):
            path = inputs / source["filename"]
            if source["file_type"] == "txt":
                write_txt(path, source, persona["folder"])
            elif source["file_type"] == "docx":
                write_docx(path, source, persona["folder"])
            elif source["file_type"] == "pdf":
                write_pdf(path, source, persona["folder"])
            else:
                raise ValueError(source["file_type"])
            source_manifest.append({"upload_order": order, "filename": f"inputs/{source['filename']}", **source["metadata"]})
            manifest_paths.append(str(path.relative_to(ROOT)))
        (folder / "source_manifest.json").write_text(json.dumps(source_manifest, indent=2) + "\n", encoding="utf-8")
        (expected_dir / "clinician_review_oracle.json").write_text(json.dumps(expanded_expected(persona["expected"]), indent=2) + "\n", encoding="utf-8")
        instructions = f"""# {persona['folder']}

{persona['summary']}

All people, organisations, events, instruments, and values in this folder are fictional. Use them only for software and clinical-workflow testing.

## Run the case

1. Create a case using `case.json`.
2. For every row in `source_manifest.json`, upload the named file and copy the listed source metadata into the upload form.
3. Wait for extraction, then review and verify appropriate evidence. Do not blindly verify every extraction.
4. Add the already-scored fictional values from instrument summary files as instrument summaries and mark them verified after checking the source.
5. Review all 18 criteria, enter a clinician-authored test conclusion, generate a draft, inspect evidence links, and test approval/rendering.
6. Compare behaviour with `expected/clinician_review_oracle.json`. Never upload the expected folder: it is a test oracle, not clinical evidence.

The suggested outcomes are deliberately not included in the upload inputs. They are expectations for testing software behaviour, not diagnoses or clinical advice.
"""
        (folder / "README.md").write_text(instructions, encoding="utf-8")
        catalog.append({"folder": persona["folder"], "cohort": persona["case"]["cohort"], "summary": persona["summary"], "source_count": len(persona["sources"])})
    (OUTPUT / "catalog.json").write_text(json.dumps(catalog, indent=2) + "\n", encoding="utf-8")
    (OUTPUT / "manifest.txt").write_text("\n".join(sorted(manifest_paths)) + "\n", encoding="utf-8")
    print(f"Generated {len(PERSONAS)} personas and {len(manifest_paths)} upload files in {OUTPUT}")


if __name__ == "__main__":
    main()
