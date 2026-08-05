from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "templates" / "adhd_report_template.docx"
FONT = "Arial Narrow"
SLATE = RGBColor(31, 60, 70)
MUTED = RGBColor(83, 110, 119)
TEAL = "6EB7C3"


def font(run, size, bold=False, italic=False, color=SLATE):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def cell_fill(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.top_margin = Inches(1.15)
section.bottom_margin = Inches(0.8)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(10)
normal.font.color.rgb = SLATE
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.08

h1 = styles["Heading 1"]
h1.font.name = FONT
h1._element.rPr.rFonts.set(qn("w:ascii"), FONT)
h1._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
h1.font.size = Pt(10.5)
h1.font.bold = True
h1.font.color.rgb = SLATE
h1.paragraph_format.space_before = Pt(9)
h1.paragraph_format.space_after = Pt(1)
h1.paragraph_format.keep_with_next = True

evidence = styles.add_style("Evidence Citation", 1)
evidence.font.name = FONT
evidence._element.rPr.rFonts.set(qn("w:ascii"), FONT)
evidence._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
evidence.font.size = Pt(7.5)
evidence.font.italic = True
evidence.font.color.rgb = MUTED
evidence.paragraph_format.left_indent = Inches(0.18)
evidence.paragraph_format.space_after = Pt(8)

header = section.header
p = header.paragraphs[0]
p.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_TAB_ALIGNMENT.RIGHT)
font(p.add_run("CLINICAL PSYCHOLOGY"), 9, bold=True, color=RGBColor.from_string("438E96"))
p.add_run("\t")
font(p.add_run("ADHD ASSESSMENT"), 9, bold=True)
p_pr = p._p.get_or_add_pPr()
borders = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "18")
bottom.set(qn("w:space"), "5")
bottom.set(qn("w:color"), TEAL)
borders.append(bottom)
p_pr.append(borders)

footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("CONFIDENTIAL • CLINICIAN-REVIEWED DRAFT  |  "), 7.5, color=MUTED)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
p._p.append(field)

# The runtime generator replaces this placeholder body completely.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("DE-IDENTIFIED ADHD ASSESSMENT REPORT TEMPLATE"), 15, bold=True)
p = doc.add_paragraph("This body is replaced by the report generator.")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.runs[0], 9, italic=True, color=MUTED)

props = doc.core_properties
props.title = "De-identified ADHD assessment report template"
props.subject = "Clinician drafting template"
props.author = "ADHD Report POC"
props.keywords = "ADHD, assessment, clinician draft"
props.comments = "Contains no patient data."

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
